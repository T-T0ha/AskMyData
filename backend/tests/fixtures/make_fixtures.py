"""Generate the deliberately messy workbooks used by the test suite.

Run with ``python -m tests.fixtures.make_fixtures`` from ``backend/``.

``messy_business.xlsx`` reproduces, in one file, every structural defect the
platform claims to handle: a title banner above the header, a two-row merged
header, currency symbols and thousands separators, six different null markers,
mixed date formats, duplicate rows, a column with no header text, and a sheet
with no primary key.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

FIXTURE_DIR = Path(__file__).resolve().parent


def _sales_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Sales Data")

    # Row 1: a title banner that must be skipped.
    ws["A1"] = "ACME Trading — Quarterly Sales Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    # Rows 2-3: a two-row header with merged group labels.
    ws["A2"] = "Order Info"
    ws.merge_cells("A2:C2")
    ws["D2"] = "Amounts"
    ws.merge_cells("D2:E2")

    ws.append([])  # placeholder so row 3 lands correctly
    for col, value in enumerate(
        ["Order ID", "Cust_ID", "Order Date", "Unit Price", "Total Amount", "Status"],
        start=1,
    ):
        ws.cell(row=3, column=col, value=value)
    ws.cell(row=2, column=6, value="Meta")

    rows = [
        ["ORD-1001", "C-001", "2024-01-15", "৳1,200.00", "৳12,000.00", "Shipped"],
        ["ORD-1002", "C-002", "16/01/2024", "$45.50", "$455.00", "Pending"],
        ["ORD-1003", "C-001", "2024-01-17", "৳2,300", "৳4,600", "shipped"],
        ["ORD-1004", "C-003", "18/01/2024", "N/A", "৳9,900", "Cancelled"],
        ["ORD-1005", "C-004", "2024-01-19", "৳780.25", "৳3,901.25", "Shipped"],
        ["ORD-1006", "C-002", "20/01/2024", "৳1,500", "৳15,000", "PENDING"],
        ["ORD-1007", "C-005", "2024-01-21", "—", "#N/A", "Shipped"],
        ["ORD-1008", "C-003", "22/01/2024", "৳3,200", "৳6,400", "Delivered"],
        ["ORD-1009", "C-001", "2024-01-23", "৳950", "৳9,500", "Delivered"],
        ["ORD-1010", "C-006", "24/01/2024", "৳1,100", "৳2,200", "Shipped"],
        # duplicate of ORD-1010 — deduplication candidate
        ["ORD-1010", "C-006", "24/01/2024", "৳1,100", "৳2,200", "Shipped"],
        ["ORD-1011", "C-004", "2024-01-25", "৳640", "৳1,280", "Pending"],
    ]
    for row in rows:
        ws.append(row)


def _customers_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Customers")
    ws.append(
        [
            "customer_id",
            "Customer Name",
            "Email Address",
            "Phone",
            "City",
            "Signup Date",
            "Active",
            "",
        ]
    )
    rows = [
        ["C-001", "Rahim Uddin", "rahim@example.com", "01711223344", "Dhaka", "01/03/2023", "Yes", None],
        ["C-002", "Karim Ali", "karim@example.com", "01822334455", "Chittagong", "15/04/2023", "Yes", None],
        ["C-003", "Nusrat Jahan", "nusrat@example.com", "01933445566", "Dhaka", "02/05/2023", "No", None],
        ["C-004", "Tanvir Hasan", "n/a", "01644556677", "Sylhet", "20/06/2023", "Yes", None],
        ["C-005", "Ayesha Siddiqua", "ayesha@example.com", "01555667788", "Dhaka", "11/07/2023", "yes", None],
        ["C-006", "Imran Khan", "imran@example.com", "?", "Khulna", "30/08/2023", "No", None],
        ["C-007", "Sadia Islam", "sadia@example.com", "01766778899", "Rajshahi", "09/09/2023", "Yes", None],
        ["C-008", "Mahmud Hossain", "mahmud@example.com", "01877889900", "Dhaka", "18/10/2023", "Yes", None],
    ]
    for row in rows:
        ws.append(row)


def _line_items_sheet(wb: Workbook) -> None:
    """A sheet with no unique key and a foreign key back to Sales Data."""

    ws = wb.create_sheet("line_items")
    ws.append(["order_id", "product_sku", "qty", "discount_pct", "line_note"])
    rows = [
        ["ORD-1001", "SKU-A1", 10, "5%", "bulk"],
        ["ORD-1001", "SKU-B2", 2, "0%", None],
        ["ORD-1002", "SKU-A1", 10, "0%", None],
        ["ORD-1003", "SKU-C3", 2, "10%", "promo"],
        ["ORD-1004", "SKU-B2", 3, "0%", None],
        ["ORD-1005", "SKU-A1", 5, "2.5%", None],
        ["ORD-1006", "SKU-D4", 10, "0%", "urgent"],
        ["ORD-1007", "SKU-C3", 1, "0%", None],
        ["ORD-1008", "SKU-B2", 2, "15%", "promo"],
        ["ORD-1009", "SKU-A1", 10, "0%", None],
        ["ORD-1010", "SKU-D4", 2, "0%", None],
        ["ORD-1011", "SKU-C3", 2, "5%", None],
    ]
    for row in rows:
        ws.append(row)


def _broken_sheet(wb: Workbook) -> None:
    """A sheet that is genuinely unusable — exercises the skip path."""

    ws = wb.create_sheet("Notes")
    ws["A1"] = "Internal notes — do not import"
    ws["A3"] = "call vendor"


def build_messy_workbook(path: Path | None = None) -> Path:
    path = path or FIXTURE_DIR / "messy_business.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    _sales_sheet(wb)
    _customers_sheet(wb)
    _line_items_sheet(wb)
    _broken_sheet(wb)
    wb.save(path)
    return path


def build_clean_workbook(path: Path | None = None) -> Path:
    """A well-formed workbook — the triage ``clean``/``fixable`` baseline."""

    path = path or FIXTURE_DIR / "clean_business.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("products")
    ws.append(["product_id", "product_name", "unit_price", "in_stock", "category"])
    rows = [
        ["P-01", "Ceiling Fan", 4500.0, 12, "Appliance"],
        ["P-02", "LED Bulb", 320.5, 340, "Lighting"],
        ["P-03", "Extension Cord", 250.0, 88, "Accessory"],
        ["P-04", "Water Pump", 12500.0, 4, "Appliance"],
        ["P-05", "Switch Board", 780.0, 150, "Accessory"],
        ["P-06", "Table Lamp", 1150.0, 26, "Lighting"],
    ]
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def build_relational_workbook(path: Path | None = None) -> Path:
    """A workbook with real structure, for Phase 3.

    ``messy_business.xlsx`` is eight to twelve rows a sheet, which is below the
    row floor foreign-key detection enforces — with that few rows, values from
    one column land inside another's by coincidence often enough that any
    result would be meaningless.  This workbook is deliberately large enough to
    be evidence, and carries one of each thing Phase 3 has to find:

    ``customers``   30 rows, keyed, with a genuine ``city → region`` hierarchy
    ``orders``      45 rows, keyed, referencing customers — with two orphans,
                    so referential integrity is *not* perfect and the negative
                    evidence sample has something real to show
    ``order_lines`` 90 rows, no single-column key, referencing orders
    ``regions``     a small lookup, below the row floor, so the "too small to
                    reason about" path is exercised on real data
    """

    path = path or FIXTURE_DIR / "relational_business.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    cities = [
        ("Dhaka", "Central"),
        ("Gazipur", "Central"),
        ("Chittagong", "South-East"),
        ("Cox's Bazar", "South-East"),
        ("Sylhet", "North-East"),
        ("Rajshahi", "North-West"),
    ]

    ws = wb.create_sheet("customers")
    ws.append(["customer_id", "customer_name", "email", "city", "region", "segment"])
    for index in range(1, 31):
        city, region = cities[index % len(cities)]
        ws.append(
            [
                f"C-{index:03d}",
                f"Customer {index}",
                f"customer{index}@example.com",
                city,
                region,
                # Deliberately on a 4-cycle while city is on a 6-cycle: on a
                # 3-cycle, city would determine segment exactly, and the
                # fixture would be asserting a coincidence of its own
                # construction rather than the city → region hierarchy.
                "Retail" if index % 4 else "Wholesale",
            ]
        )

    ws = wb.create_sheet("orders")
    ws.append(["order_id", "customer_id", "order_date", "status", "total_amount"])
    for index in range(1, 46):
        # Two orders reference a customer that is not in the customers sheet:
        # the orphan case the negative evidence panel exists to surface.
        customer = f"C-{999 if index in (7, 23) else (index % 30) + 1:03d}"
        ws.append(
            [
                f"ORD-{2000 + index}",
                customer,
                f"2024-{(index % 12) + 1:02d}-{(index % 27) + 1:02d}",
                ["Shipped", "Pending", "Delivered"][index % 3],
                round(500 + index * 37.5, 2),
            ]
        )

    ws = wb.create_sheet("order_lines")
    ws.append(["order_id", "sku", "quantity", "unit_price"])
    for index in range(1, 91):
        ws.append(
            [
                f"ORD-{2000 + (index % 45) + 1}",
                f"SKU-{(index % 12) + 1:02d}",
                (index % 7) + 1,
                round(25 + (index % 9) * 12.5, 2),
            ]
        )

    ws = wb.create_sheet("regions")
    # Two numeric columns, not one.  Phase 0 calls a row a header when over
    # 60 % of its cells are non-numeric, so a lookup of "text, text, number"
    # has its first data row absorbed into the header — the ambiguity the
    # README documents.  Half the row being numeric settles it, and a fixture
    # should exercise the path it is aimed at rather than trip over another.
    ws.append(["region", "manager", "target_revenue", "headcount"])
    for index, region in enumerate(sorted({region for _, region in cities}), start=1):
        ws.append([region, f"Manager of {region}", 100000 * index, 4 + index])

    wb.save(path)
    return path


if __name__ == "__main__":
    print(build_messy_workbook())
    print(build_clean_workbook())
    print(build_relational_workbook())
